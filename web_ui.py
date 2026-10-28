#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Уеб интерфейс за преглед на събраните обществени поръчки (ЦАИС ЕОП)
=====================================================================

Показва данните от SQLite базата `eop_history.db` (създавана от
`collect_history.py`) в табличен вид на адрес:

    http://127.0.0.1:7000/

Възможности:
  * поле за търсене;
  * сортиране по колони: дата, име/възложител,
    начин на възлагане, обект, лице за контакт, статус, стойност и др.;
  * динамично разгъване на ред — клик върху поръчка показва наведнъж
    ЦЯЛАТА информация (всички полета + пълен списък на участниците);
  * статистически данни;
  
Търсенето, сортирането и страницирането се изпълняват сървърно (SQL),
за да работи бързо с ~136 000 записа.

Употреба:
    python3 web_ui.py [--port 7000] [--db eop_history.db]
"""

import argparse
import configparser
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs

import httpx

DEFAULT_HOST = "127.0.0.1"
DEFAULT_PORT = 7000
DEFAULT_DB = str(Path(__file__).with_name("eop_history.db"))

# Конфигурация за DeepSeek (същото име като скрипта, в същата папка).
CONFIG_PATH = str(Path(__file__).with_name("web_ui.cfg"))

# Услуга на ЦАИС ЕОП (за извличане на прикачените файлове по заявка).
EOP_SERVICE = "https://service.eop.bg/NX1Service.svc"

# Допустими колони за сортиране (ключ от UI -> SQL израз). Whitelist.
SORT_COLUMNS = {
    "special_number": "special_number",
    "name": "name",
    "organization": "organization",
    "procedure": "procedure_type_label",
    "object": "type_of_contract_label",
    "value": "estimated_value",
    "date": "publication_date",
    "status": "COALESCE(NULLIF(outcome, ''), status)",
    "contact": "contact_person",
    "participants": "participants_count",
}

# Колони, в които се търси (LIKE).
SEARCH_COLUMNS = [
    "special_number", "name", "organization", "procedure_type_label",
    "type_of_contract_label", "currency", "contact_person", "status",
    "outcome", "parent_tender_special_number", "participants",
]

# Фиксиран курс за нормализиране на прогнозните стойности (BGN -> EUR).
BGN_TO_EUR = 1.95583

# Имена на месеците за статистиката по месеци.
MONTH_NAMES = {
    1: "Януари", 2: "Февруари", 3: "Март", 4: "Април", 5: "Май", 6: "Юни",
    7: "Юли", 8: "Август", 9: "Септември", 10: "Октомври", 11: "Ноември", 12: "Декември",
}

# Ключови думи за разпознаване на града от името на възложителя.
CITY_KEYWORDS = [
    ("София", ["софия", "софийск", "столична"]),
    ("Пловдив", ["пловдив"]),
    ("Варна", ["варна"]),
    ("Бургас", ["бургас"]),
    ("Русе", ["русе"]),
    ("Стара Загора", ["стара загора", "старозагор"]),
    ("Плевен", ["плевен"]),
    ("Сливен", ["сливен"]),
    ("Добрич", ["добрич"]),
    ("Шумен", ["шумен"]),
    ("Перник", ["перник"]),
    ("Хасково", ["хасково"]),
    ("Ямбол", ["ямбол"]),
    ("Пазарджик", ["пазарджик"]),
    ("Благоевград", ["благоевград"]),
    ("Велико Търново", ["велико търново", "великотърнов"]),
    ("Враца", ["враца"]),
    ("Габрово", ["габрово"]),
    ("Кюстендил", ["кюстендил"]),
    ("Видин", ["видин"]),
    ("Монтана", ["монтана"]),
    ("Ловеч", ["ловеч"]),
    ("Разград", ["разград"]),
    ("Силистра", ["силистра"]),
    ("Търговище", ["търговище"]),
    ("Кърджали", ["кърджали"]),
    ("Смолян", ["смолян"]),
    ("Казанлък", ["казанлък"]),
    ("Козлодуй", ["козлодуй"]),
    ("Асеновград", ["асеновград"]),
    ("Димитровград", ["димитровград"]),
    ("Самоков", ["самоков"]),
    ("Сандански", ["сандански"]),
    ("Петрич", ["петрич"]),
    ("Севлиево", ["севлиево"]),
    ("Карлово", ["карлово"]),
    ("Троян", ["троян"]),
    ("Дупница", ["дупница"]),
    ("Горна Оряховица", ["горна оряховица"]),
    ("Свищов", ["свищов"]),
    ("Велинград", ["велинград"]),
    ("Ботевград", ["ботевград"]),
    ("Поморие", ["поморие"]),
    ("Несебър", ["несебър"]),
    ("Балчик", ["балчик"]),
    ("Нова Загора", ["нова загора"]),
    ("Радомир", ["радомир"]),
    ("Ихтиман", ["ихтиман"]),
    ("Сопот", ["сопот"]),
    ("Гоце Делчев", ["гоце делчев"]),
    ("Разлог", ["разлог"]),
    ("Панагюрище", ["панагюрище"]),
]


def extract_city(org_name):
    """Разпознава града от името на възложителя (best-effort)."""
    if not org_name:
        return "Други"
    low = org_name.lower()
    for city, keys in CITY_KEYWORDS:
        for key in keys:
            if key in low:
                return city
    return "Други"


def load_page(db_path, query, sort_key, order, page, size):
    """Връща dict {total, rows} за една страница от резултатите."""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    # Вградената LOWER() на SQLite не обработва кирилица. Регистрираме
    # собствена функция, за да търсим без значение на малки/главни букви.
    conn.create_function("lower_cyr", 1, lambda s: (s or "").lower())

    where = ""
    params = []
    q = (query or "").strip()
    if q:
        like = f"%{q.lower()}%"
        clauses = [f"lower_cyr({col}) LIKE ?" for col in SEARCH_COLUMNS]
        where = "WHERE " + " OR ".join(clauses)
        params = [like] * len(SEARCH_COLUMNS)

    total = conn.execute(f"SELECT COUNT(*) FROM tenders {where}", params).fetchone()[0]

    sort_expr = SORT_COLUMNS.get(sort_key, "publication_date")
    order_sql = "ASC" if order == "asc" else "DESC"
    # Стабилно вторично сортиране по уникален номер.
    sql = (
        f"SELECT * FROM tenders {where} "
        f"ORDER BY {sort_expr} {order_sql}, special_number {order_sql} "
        f"LIMIT ? OFFSET ?"
    )
    rows = conn.execute(sql, params + [size, (page - 1) * size]).fetchall()
    conn.close()

    out = []
    for r in rows:
        d = dict(r)
        for col in ("participants", "specialty", "publication_form_types"):
            try:
                d[col] = json.loads(d.get(col) or "[]")
            except Exception:
                d[col] = [] if col != "specialty" else {}
        out.append(d)
    return {"total": total, "rows": out}


def load_stats(db_path):
    """Изчислява агрегираните статистики за прегледа."""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        top_authorities = [dict(r) for r in conn.execute(
            "SELECT organization AS name, COUNT(*) AS cnt FROM tenders "
            "WHERE organization IS NOT NULL AND organization != '' "
            "GROUP BY organization ORDER BY cnt DESC, organization LIMIT 10"
        ).fetchall()]

        statuses = [dict(r) for r in conn.execute(
            "SELECT COALESCE(NULLIF(outcome, ''), NULLIF(status, ''), 'Без статус') AS label, "
            "COUNT(*) AS cnt FROM tenders GROUP BY label ORDER BY cnt DESC"
        ).fetchall()]

        all_values = conn.execute(
            "SELECT special_number, name, organization, estimated_value, currency "
            "FROM tenders WHERE estimated_value IS NOT NULL"
        ).fetchall()
        vals = []
        for r in all_values:
            cur = r["currency"]
            eur = r["estimated_value"] if cur == "EUR" else r["estimated_value"] / BGN_TO_EUR
            vals.append({
                "special_number": r["special_number"],
                "name": r["name"],
                "organization": r["organization"],
                "estimated_value": r["estimated_value"],
                "currency": cur,
                "estimated_value_eur": round(eur, 2),
            })
        vals.sort(key=lambda x: -x["estimated_value_eur"])
        top_expensive = vals[:10]

        top_participants = [dict(r) for r in conn.execute(
            "SELECT special_number, name, organization, participants_count FROM tenders "
            "WHERE participants_count IS NOT NULL AND participants_count > 0 "
            "ORDER BY participants_count DESC, special_number LIMIT 10"
        ).fetchall()]

        # Статистика по месеци (Януари - Декември).
        month_rows = conn.execute(
            "SELECT CAST(substr(publication_date, 6, 2) AS INTEGER) AS m, COUNT(*) AS cnt "
            "FROM tenders WHERE publication_date IS NOT NULL AND publication_date != '' "
            "GROUP BY m"
        ).fetchall()
        month_map = {r["m"]: r["cnt"] for r in month_rows if r["m"] is not None}
        monthly = [{"month": i, "label": MONTH_NAMES[i], "cnt": month_map.get(i, 0)} for i in range(1, 13)]

        # Статистика по години (в хронологичен ред).
        yearly = [dict(r) for r in conn.execute(
            "SELECT CAST(substr(publication_date, 1, 4) AS INTEGER) AS year, COUNT(*) AS cnt "
            "FROM tenders WHERE publication_date IS NOT NULL AND publication_date != '' "
            "GROUP BY year ORDER BY year"
        ).fetchall()]

        # Стойност на поръчките по години (нормализирана в EUR) + обща сума.
        value_rows = conn.execute(
            "SELECT CAST(substr(publication_date, 1, 4) AS INTEGER) AS year, COUNT(*) AS cnt, "
            "SUM(CASE WHEN currency = 'EUR' THEN estimated_value ELSE estimated_value / ? END) AS total_eur "
            "FROM tenders "
            "WHERE publication_date IS NOT NULL AND publication_date != '' AND estimated_value IS NOT NULL "
            "GROUP BY year ORDER BY year",
            (BGN_TO_EUR,),
        ).fetchall()
        yearly_values = [
            {"year": r["year"], "cnt": r["cnt"], "total_eur": round(r["total_eur"] or 0.0, 2)}
            for r in value_rows if r["year"] is not None
        ]
        total_value_eur = round(sum(v["total_eur"] for v in yearly_values), 2)

        # Топ 20 градове (по името на възложителя).
        org_rows = conn.execute(
            "SELECT organization FROM tenders WHERE organization IS NOT NULL AND organization != ''"
        ).fetchall()
        city_counts = {}
        for r in org_rows:
            city = extract_city(r["organization"])
            city_counts[city] = city_counts.get(city, 0) + 1
        top_cities = [{"city": c, "cnt": n} for c, n in sorted(city_counts.items(), key=lambda x: (-x[1], x[0]))[:20]]

        total = conn.execute("SELECT COUNT(*) FROM tenders").fetchone()[0]
    finally:
        conn.close()

    return {
        "total": total,
        "top_authorities": top_authorities,
        "statuses": statuses,
        "top_expensive": top_expensive,
        "top_participants": top_participants,
        "monthly": monthly,
        "yearly": yearly,
        "top_cities": top_cities,
        "yearly_values": yearly_values,
        "total_value_eur": total_value_eur,
    }


# ---------------------------------------------------------------------------
# Анализ на поръчка с DeepSeek
# ---------------------------------------------------------------------------

MAX_FILES = 10                 # максимален брой прикачени файлове, чиито линкове подаваме
MAX_FILE_BYTES = 10 * 1024 * 1024   # пропускаме файлове по-големи от 10 MB
MAX_TEXT_PER_FILE = 15000      # максимум символи текст от един файл
MAX_TOTAL_TEXT = 40000         # общ максимум текст, подаван към модела

# Точките от анализа — всяка се извиква в отделна заявка към AI, за да не се
# изрязва отговорът (ограничение в max_tokens) и да има обобщение за всяка точка.
ANALYSIS_POINTS = [
    "Цел на поръчката",
    "Обхват на поръчката",
    "Специфични изисквания",
    "Изисквана квалификация/експертиза от участниците",
    "Има ли доставка на оборудване",
    "Период/време за изпълнение и краен срок за подаване на оферти",
]

POINT_PROMPT_TEMPLATE = """Ти си експерт по обществени поръчки. Анализирай документацията на следната обществена поръчка.

URL на поръчката (ЦАИС ЕОП): {url}

Прикачени файлове (линкове за справка):
{files_list}

Извлечено съдържание на прикачените файлове:
{documents_text}

Отговори на български език САМО на следната точка от анализа:
{point}

Правила:
- Опиши директно намереното по тази точка, без встъпителни фрази и без уводни думи като „Да, има...“ или „Не, няма...“.
- Работи само по тази точка и по съдържанието на документите, без допълнителна информация или разяснения.
- Дай кратко обобщение на анализираните документи по тази точка."""


def load_deepseek_config():
    """Чете настройките за DeepSeek от web_ui.cfg (до скрипта)."""
    if not Path(CONFIG_PATH).exists():
        raise RuntimeError(
            f"Липсва конфигурационен файл {CONFIG_PATH}. "
            "Създайте го с раздел [deepseek] и ключове: api_key, base_url, model, timeout."
        )
    cfg = configparser.ConfigParser()
    cfg.read(CONFIG_PATH, encoding="utf-8")
    if "deepseek" not in cfg:
        raise RuntimeError("Липсва раздел [deepseek] в конфигурационния файл.")
    sec = cfg["deepseek"]
    api_key = sec.get("api_key", "").strip()
    if not api_key or api_key in ("YOUR_API_KEY", "sk-xxxx", "PASTE_YOUR_KEY"):
        raise RuntimeError("Не е зададен валиден api_key за DeepSeek в конфигурационния файл.")
    return {
        "api_key": api_key,
        "base_url": sec.get("base_url", "https://api.deepseek.com").strip().rstrip("/"),
        "model": sec.get("model", "deepseek-chat").strip(),
        "timeout": sec.getfloat("timeout", 300.0),
    }


def load_notifications_config():
    """Чете настройките за известяване по имейл от web_ui.cfg."""
    defaults = {
        "enabled": False,
        "keywords": "",
        "organization": "",
        "object": "",
        "emails": "",
    }
    if not Path(CONFIG_PATH).exists():
        return defaults
    cfg = configparser.ConfigParser(interpolation=None)
    cfg.read(CONFIG_PATH, encoding="utf-8")
    if "notifications" not in cfg:
        return defaults
    sec = cfg["notifications"]
    return {
        "enabled": sec.getboolean("enabled", False),
        "keywords": sec.get("keywords", "").strip(),
        "organization": sec.get("organization", "").strip(),
        "object": sec.get("object", "").strip(),
        "emails": sec.get("emails", "").strip(),
    }


def load_smtp_config():
    """Чете SMTP настройките от web_ui.cfg."""
    defaults = {"host": "", "port": 587, "username": "", "password": "", "from": "", "use_tls": True}
    if not Path(CONFIG_PATH).exists():
        return defaults
    cfg = configparser.ConfigParser(interpolation=None)
    cfg.read(CONFIG_PATH, encoding="utf-8")
    if "smtp" not in cfg:
        return defaults
    sec = cfg["smtp"]
    return {
        "host": sec.get("host", "").strip(),
        "port": sec.getint("port", 587),
        "username": sec.get("username", "").strip(),
        "password": sec.get("password", "").strip(),
        "from": sec.get("from", "").strip(),
        "use_tls": sec.getboolean("use_tls", True),
    }


def _update_config_sections(sections):
    """Записва секции в web_ui.cfg, запазвайки останалите секции и коментари."""
    lines = []
    if Path(CONFIG_PATH).exists():
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            lines = f.read().splitlines()

    out = []
    i = 0
    replaced = set()
    while i < len(lines):
        stripped = lines[i].strip()
        if stripped.startswith("[") and stripped.endswith("]"):
            name = stripped[1:-1].strip()
            if name in sections:
                out.extend(sections[name])
                replaced.add(name)
                i += 1
                while i < len(lines):
                    s = lines[i].strip()
                    if s.startswith("[") and s.endswith("]"):
                        break
                    i += 1
                continue
        out.append(lines[i])
        i += 1

    for name, content in sections.items():
        if name not in replaced:
            if out and out[-1].strip() != "":
                out.append("")
            out.extend(content)

    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")


def save_notifications_config(data):
    """Записва секциите [notifications] и [smtp] в web_ui.cfg."""
    notif = [
        "[notifications]",
        f"enabled = {'true' if data.get('enabled') else 'false'}",
        f"keywords = {data.get('keywords', '') or ''}",
        f"organization = {data.get('organization', '') or ''}",
        f"object = {data.get('object', '') or ''}",
        f"emails = {data.get('emails', '') or ''}",
    ]
    smtp_data = data.get("smtp") or {}
    smtp = [
        "[smtp]",
        f"host = {smtp_data.get('host', '') or ''}",
        f"port = {smtp_data.get('port', 587) or 587}",
        f"username = {smtp_data.get('username', '') or ''}",
        f"password = {smtp_data.get('password', '') or ''}",
        f"from = {smtp_data.get('from', '') or ''}",
        f"use_tls = {'true' if smtp_data.get('use_tls') else 'false'}",
    ]
    _update_config_sections({"notifications": notif, "smtp": smtp})


def fetch_attachments(tender_id):
    """Взима списъка с прикачените файлове и подписани URL за сваляне."""
    attachments = []
    with httpx.Client(timeout=60, verify=False, follow_redirects=True) as client:
        r = client.post(
            EOP_SERVICE + "/GetPublishedTenderDetails",
            json={"tenderId": tender_id, "ianaTimeZone": "Europe/Sofia"},
        )
        if r.status_code != 200:
            return attachments
        details = r.json()
        docs = details.get("TenderDescriptionDocuments") or []
        for doc in docs[:MAX_FILES]:
            did = doc.get("Id")
            name = doc.get("Name") or "(без име)"
            ext = (doc.get("Extension") or "").lower()
            size = doc.get("Size") or 0
            url = ""
            if did:
                try:
                    rr = client.post(
                        EOP_SERVICE + "/GetSignedUrlByDocumentId",
                        json={"documentId": did},
                    )
                    if rr.status_code == 200:
                        url = rr.json().get("Url") or ""
                except Exception:
                    url = ""
            attachments.append({"name": name, "extension": ext, "size": size, "url": url})
    return attachments


def _decode_text(data):
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            text = data.decode(enc)
            printable = sum(1 for ch in text if ch.isprintable() or ch in "\n\r\t")
            if printable / max(len(text), 1) > 0.9:
                return text
        except Exception:
            continue
    return ""


def _extract_pdf(data):
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                parts.append("")
        return "\n\n".join(parts)
    except Exception:
        return ""


def _extract_pdftotext(data):
    fd, path = tempfile.mkstemp(suffix=".pdf")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        out = subprocess.run(["pdftotext", path, "-"], capture_output=True, timeout=60)
        return out.stdout.decode("utf-8", errors="replace")
    except Exception:
        return ""
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def _extract_docx(data):
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for tbl in doc.tables:
            for row in tbl.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _extract_via_soffice(data, filename):
    tmpdir = tempfile.mkdtemp()
    try:
        src = os.path.join(tmpdir, filename)
        with open(src, "wb") as f:
            f.write(data)
        subprocess.run(
            [
                "soffice", "--headless",
                f"-env:UserInstallation=file://{tmpdir}/lo_profile",
                "--convert-to", "txt:Text (encoded):UTF8",
                "--outdir", tmpdir, src,
            ],
            capture_output=True, timeout=180,
        )
        out_path = os.path.join(tmpdir, Path(filename).stem + ".txt")
        if os.path.exists(out_path):
            with open(out_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        return ""
    except Exception:
        return ""
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def _extract_printable_runs(data):
    runs = re.findall(rb"[ -~\xC0-\xFF\n\r\t]{4,}", data)
    text = b"\n".join(runs)
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            return text.decode(enc)
        except Exception:
            continue
    return text.decode("latin-1", errors="replace")


def _extract_zip(data, depth=1):
    if depth > 2:
        return ""
    parts = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                if info.is_dir():
                    continue
                try:
                    inner = z.read(info.filename)
                except Exception:
                    continue
                parts.append(f"--- {info.filename} ---")
                parts.append(extract_text(inner, info.filename))
    except Exception:
        return ""
    return "\n".join(parts)


def extract_text(data, filename):
    """Извлича текст от файл: txt/pdf/docx/doc/xls/zip и др. (best-effort)."""
    if not data:
        return ""
    ext = Path(filename).suffix.lower()

    if ext in (".txt", ".csv", ".md", ".xml", ".html", ".htm", ".json", ".log", ".ini", ".rtf"):
        text = _decode_text(data)
        if text:
            return text

    if ext == ".pdf":
        text = _extract_pdf(data).strip()
        if len(text) > 40:
            return text
        text = _extract_pdftotext(data).strip()
        if text:
            return text

    if ext == ".docx":
        text = _extract_docx(data).strip()
        if text:
            return text
        return _extract_via_soffice(data, filename)

    if ext in (".doc", ".xls", ".xlsx", ".odt", ".ods"):
        return _extract_via_soffice(data, filename)

    if ext == ".zip":
        return _extract_zip(data)

    return _extract_printable_runs(data)


def call_deepseek(messages, config, max_tokens=2000):
    """Извиква DeepSeek (OpenAI-съвместим chat/completions) и връща текста."""
    url = config["base_url"] + "/chat/completions"
    headers = {
        "Authorization": f"Bearer {config['api_key']}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": config["model"],
        "messages": messages,
        "temperature": 0.2,
        "max_tokens": max_tokens,
        "stream": False,
    }
    with httpx.Client(timeout=config["timeout"], verify=False) as client:
        r = client.post(url, headers=headers, json=payload)
    if r.status_code != 200:
        raise RuntimeError(f"DeepSeek грешка {r.status_code}: {r.text[:300]}")
    data = r.json()
    try:
        return data["choices"][0]["message"]["content"]
    except (KeyError, IndexError):
        raise RuntimeError(f"Неочакван отговор от DeepSeek: {json.dumps(data)[:300]}")


def analyze_tender_stream(tender_id):
    """Генератор: извлича файловете и връща AI анализ със събития за напредък.

    Всяко събитие е dict: {"type": "progress"|"done", ...}.
    """
    yield {"type": "progress", "message": "Зареждане на списъка с прикачени файлове..."}
    config = load_deepseek_config()
    url = f"https://app.eop.bg/today/{tender_id}"
    attachments = fetch_attachments(tender_id)
    total = len(attachments)

    files_lines = []
    doc_parts = []
    total_text = 0
    with httpx.Client(timeout=120, verify=False, follow_redirects=True) as client:
        for idx, att in enumerate(attachments, 1):
            files_lines.append(
                f"- {att['name']} ({att['extension'] or 'без разширение'}) -> {att['url'] or '(няма линк)'}"
            )
            yield {"type": "progress", "message": f"Изтегляне и обработка на файл {idx}/{total}: {att['name']}"}
            if not att["url"] or att["size"] > MAX_FILE_BYTES:
                doc_parts.append(f"=== {att['name']} ===\n(файлът е пропуснат: твърде голям или без линк)")
                continue
            try:
                resp = client.get(att["url"])
                if resp.status_code != 200:
                    doc_parts.append(f"=== {att['name']} ===\n(грешка при изтегляне: HTTP {resp.status_code})")
                    continue
                text = (extract_text(resp.content, att["name"]) or "").strip()
                text = re.sub(r"[ \t]+", " ", text)
                if not text:
                    doc_parts.append(f"=== {att['name']} ===\n(съдържанието не можа да бъде извлечено)")
                    continue
                text = text[:MAX_TEXT_PER_FILE]
                if total_text + len(text) > MAX_TOTAL_TEXT:
                    text = text[:max(0, MAX_TOTAL_TEXT - total_text)]
                doc_parts.append(f"=== {att['name']} ===\n{text}")
                total_text += len(text)
            except Exception as exc:
                doc_parts.append(f"=== {att['name']} ===\n(грешка при изтегляне: {exc})")

    documents_text = "\n\n".join(doc_parts) if doc_parts else "(не е налично извлечено съдържание)"
    files_list = "\n".join(files_lines) if files_lines else "(няма прикачени файлове)"

    results = []
    for i, point in enumerate(ANALYSIS_POINTS, 1):
        yield {"type": "progress", "message": f"Анализиране на точка {i}/{len(ANALYSIS_POINTS)}: {point}..."}
        prompt = POINT_PROMPT_TEMPLATE.format(
            url=url,
            files_list=files_list,
            documents_text=documents_text,
            point=point,
        )
        answer = call_deepseek(
            [
                {"role": "system", "content": "Ти си експерт по обществени поръчки. Анализираш предоставеното извлечено съдържание на тръжната документация и отговаряш директно по точка, без встъпителни фрази и без допълнителна информация."},
                {"role": "user", "content": prompt},
            ],
            config,
        )
        results.append(f"{i}. {point}:\n{answer.strip()}")

    combined = "\n\n".join(results)
    yield {"type": "done", "url": url, "analysis": combined, "files": total}


HTML_TEMPLATE = r"""<!DOCTYPE html>
<html lang="bg">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>ЦАИС ЕОП — Регистър на обществените поръчки</title>
<style>
  :root {
    --bg: #f2f4f7; --panel: #fff; --border: #e3e8ef; --text: #1f2937;
    --muted: #6b7280; --accent: #003a51; --accent2: #189cff; --hover: #f6f9fc;
  }
  * { box-sizing: border-box; }
  body { margin: 0; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif; background: var(--bg); color: var(--text); }
  header { background: var(--accent); color: #fff; padding: 14px 24px; }
  header .top { display: flex; align-items: center; gap: 16px; flex-wrap: wrap; }
  header h1 { font-size: 18px; margin: 0; font-weight: 600; }
  header .meta { margin-left: auto; font-size: 13px; opacity: .95; display: flex; gap: 12px; align-items: center; }
  .searchbox { display: flex; align-items: center; background: #fff; border-radius: 8px; padding: 0 10px; gap: 8px; flex: 1 1 320px; max-width: 560px; }
  .searchbox input { border: 0; outline: 0; padding: 9px 0; font-size: 14px; width: 100%; background: transparent; }
  .searchbox .icon { color: var(--muted); }
  .searchbox input::placeholder { color: #9aa4b2; }
  button { cursor: pointer; border: 0; border-radius: 6px; }
  .refresh { background: var(--accent2); color: #fff; padding: 8px 14px; font-size: 13px; }
  .refresh:hover { filter: brightness(1.08); }
  .notif-btn { background: rgba(255,255,255,.14); color: #fff; padding: 8px 14px; font-size: 13px; border: 1px solid rgba(255,255,255,.25); }
  .notif-btn:hover { background: rgba(255,255,255,.25); }
  .wrap { padding: 16px 24px 40px; }
  .toolbar { display: flex; align-items: center; gap: 12px; margin-bottom: 12px; flex-wrap: wrap; }
  .toolbar .stats { color: var(--muted); font-size: 13px; }
  .pager { margin-left: auto; display: flex; align-items: center; gap: 8px; }
  .pager button { background: #fff; border: 1px solid var(--border); padding: 6px 12px; font-size: 13px; }
  .pager button:disabled { opacity: .4; cursor: default; }
  .pager select { padding: 6px; border: 1px solid var(--border); border-radius: 6px; background: #fff; }
  .table-scroll { overflow: auto; background: var(--panel); border: 1px solid var(--border); border-radius: 10px; box-shadow: 0 1px 3px rgba(16,24,40,.06); }
  table { border-collapse: separate; border-spacing: 0; font-size: 13px; table-layout: fixed; }
  thead th { position: sticky; top: 0; z-index: 3; background: #f8fafc; border-bottom: 1px solid var(--border); border-right: 1px solid var(--border); padding: 0; text-align: left; user-select: none; }
  thead th .th-inner { display: flex; align-items: center; gap: 6px; padding: 10px 10px; white-space: nowrap; overflow: hidden; }
  thead th.sortable { cursor: pointer; }
  thead th.sortable:hover { background: #eef3f8; color: var(--accent); }
  thead th .label { overflow: hidden; text-overflow: ellipsis; }
  thead th .arrow { opacity: .4; font-size: 11px; flex-shrink: 0; }
  thead th.sorted .arrow { opacity: 1; color: var(--accent2); }
  .resizer { position: absolute; top: 0; right: 0; width: 8px; height: 100%; cursor: col-resize; z-index: 5; }
  .resizer:hover { background: rgba(24,156,255,.35); }
  tbody td { padding: 8px 10px; border-bottom: 1px solid var(--border); border-right: 1px solid var(--border); vertical-align: top; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
  tbody tr.main { cursor: pointer; }
  tbody tr.main:hover { background: var(--hover); }
  tbody tr.main.open { background: #eaf4fc; }
  tbody tr.detail td { white-space: normal; background: #fbfdff; padding: 16px; }
  .num { font-weight: 600; color: var(--accent); }
  .muted { color: var(--muted); }
  .badge { display: inline-block; border-radius: 999px; padding: 2px 10px; font-size: 12px; font-weight: 600; }
  .badge.ok { background: #e6f6ec; color: #1a7f37; }
  .badge.term { background: #fdecec; color: #b42318; }
  .badge.award { background: #eef3ff; color: #2452c6; }
  .badge.neutral { background: #eef3f8; color: var(--accent); }
  .empty { padding: 40px; text-align: center; color: var(--muted); }
  .detail-grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(300px, 1fr)); gap: 10px 22px; }
  .field { min-width: 0; }
  .field .k { font-size: 11px; color: var(--muted); text-transform: uppercase; letter-spacing: .3px; }
  .field .v { font-size: 13px; word-break: break-word; }
  .detail-title { font-weight: 600; margin: 14px 0 6px; font-size: 14px; color: var(--accent); }
  .detail-title:first-child { margin-top: 0; }
  table.participants { width: 100%; table-layout: auto; border: 1px solid var(--border); border-radius: 8px; overflow: hidden; font-size: 12px; }
  table.participants th, table.participants td { padding: 6px 10px; border-bottom: 1px solid var(--border); white-space: normal; }
  table.participants th { background: #f1f5f9; text-align: left; font-weight: 600; }
  .chips { display: flex; flex-wrap: wrap; gap: 6px; }
  .chip { background: #eef3f8; border-radius: 6px; padding: 3px 8px; font-size: 12px; }
  a.link { color: var(--accent2); text-decoration: none; }
  a.link:hover { text-decoration: underline; }
  .papagal-link { color: var(--accent2); text-decoration: none; font-size: 14px; }
  .papagal-link:hover { filter: brightness(1.15); }
  .searchbox button { background: var(--accent2); color: #fff; border-radius: 6px; padding: 6px 14px; font-size: 13px; white-space: nowrap; }
  .searchbox button:hover { filter: brightness(1.08); }
  .pageinput { width: 64px; padding: 6px; border: 1px solid var(--border); border-radius: 6px; background: #fff; text-align: center; }
  .pager .goto { background: #fff; border: 1px solid var(--border); padding: 6px 12px; font-size: 13px; }
  .pager .goto:hover { background: #eef3f8; }
  .modal { position: fixed; inset: 0; background: rgba(15,23,42,.45); display: flex; align-items: center; justify-content: center; z-index: 1000; }
  .modal-content { background: #fff; border-radius: 12px; padding: 28px 34px; display: flex; flex-direction: column; align-items: center; gap: 14px; box-shadow: 0 10px 30px rgba(0,0,0,.2); font-size: 14px; color: var(--text); }
  .spinner { width: 34px; height: 34px; border: 4px solid #e3e8ef; border-top-color: var(--accent2); border-radius: 50%; animation: spin .7s linear infinite; }
  @keyframes spin { to { transform: rotate(360deg); } }
  .tabs { display: flex; gap: 6px; }
  .tab { background: rgba(255,255,255,.12); color: #fff; border: 0; border-radius: 8px; padding: 8px 16px; font-size: 14px; }
  .tab:hover { background: rgba(255,255,255,.22); }
  .tab.active { background: #fff; color: var(--accent); font-weight: 600; }
  .stats-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 18px; }
  @media (max-width: 900px) { .stats-grid { grid-template-columns: 1fr; } }
  .stat-card { background: var(--panel); border: 1px solid var(--border); border-radius: 12px; padding: 18px 20px; box-shadow: 0 1px 3px rgba(16,24,40,.06); }
  .stat-card h3 { margin: 0 0 14px; font-size: 15px; color: var(--accent); }
  .bar-row { display: flex; align-items: center; gap: 10px; margin-bottom: 8px; font-size: 12px; }
  .bar-label { flex: 0 0 200px; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
  .bar-track { flex: 1; background: #eef2f7; border-radius: 6px; height: 16px; overflow: hidden; }
  .bar-fill { height: 100%; background: linear-gradient(90deg, #189cff, #6cc4ff); border-radius: 6px; }
  .bar-value { flex: 0 0 40px; text-align: right; font-weight: 600; }
  .pie-wrap { display: flex; align-items: center; gap: 20px; flex-wrap: wrap; }
  .legend { border-collapse: collapse; font-size: 12px; }
  .legend td { padding: 3px 8px; }
  .legend .dot { display: inline-block; width: 10px; height: 10px; border-radius: 3px; margin-right: 6px; }
  .minitable-scroll { overflow-x: auto; }
  .minitable { border-collapse: separate; border-spacing: 0; font-size: 12px; table-layout: fixed; }
  .minitable th { position: relative; background: #f1f5f9; font-weight: 600; padding: 6px 8px; border-bottom: 1px solid var(--border); border-right: 1px solid var(--border); text-align: left; }
  .minitable th .th-inner { white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
  .minitable td { padding: 6px 8px; border-bottom: 1px solid var(--border); border-right: 1px solid var(--border); text-align: left; vertical-align: top; white-space: normal; word-break: break-word; overflow-wrap: anywhere; }
  .minitable .muted { font-size: 11px; }
  .stats-note { color: var(--muted); font-size: 12px; margin-top: 8px; }
  .analyze-btn { background: var(--accent2); color: #fff; border: 0; border-radius: 6px; padding: 9px 18px; font-size: 13px; cursor: pointer; margin-top: 14px; }
  .analyze-btn:hover { filter: brightness(1.08); }
  .analysis-content { width: min(760px, 92vw); max-height: 82vh; display: flex; flex-direction: column; align-items: stretch; padding: 20px 22px; }
  .analysis-header { display: flex; align-items: center; justify-content: space-between; gap: 12px; font-weight: 600; font-size: 15px; }
  .analysis-header .close { background: none; border: 0; font-size: 20px; cursor: pointer; color: var(--muted); line-height: 1; }
  .analysis-header .close:hover { color: var(--text); }
  .analysis-body { overflow: auto; white-space: pre-wrap; font-size: 13px; line-height: 1.55; margin-top: 12px; flex: 1 1 auto; }
  .analysis-actions { display: flex; justify-content: flex-end; gap: 10px; margin-top: 16px; }
  .analysis-actions button { padding: 8px 16px; border-radius: 6px; font-size: 13px; }
  #analysis-cancel { background: #fff; border: 1px solid var(--border); color: var(--text); }
  #analysis-cancel:hover { background: #f6f8fa; }
  #analysis-ok { background: var(--accent2); border: 0; color: #fff; }
  .notif-content { width: min(620px, 94vw); max-height: 85vh; overflow: auto; display: flex; flex-direction: column; gap: 12px; padding: 20px 24px; }
  .notif-content h3 { margin: 0; font-size: 15px; }
  .form-row { display: flex; flex-direction: column; gap: 5px; }
  .form-row label { font-size: 12px; color: var(--muted); font-weight: 600; }
  .form-row input[type="text"], .form-row select { padding: 8px 10px; border: 1px solid var(--border); border-radius: 6px; font-size: 13px; width: 100%; box-sizing: border-box; }
  .form-row .hint { font-size: 11px; color: var(--muted); }
  .form-check { display: flex; align-items: center; gap: 8px; font-size: 13px; }
  .form-check input { width: auto; }
  .form-actions { display: flex; justify-content: flex-end; gap: 10px; margin-top: 4px; }
  .form-actions button { padding: 8px 16px; border-radius: 6px; font-size: 13px; }
  #notif-save { background: var(--accent2); border: 0; color: #fff; }
  #notif-cancel { background: #fff; border: 1px solid var(--border); }
  #notif-status { font-size: 12px; color: var(--muted); min-height: 16px; }
</style>
</head>
<body>
<header>
  <div class="top">
    <h1>ЦАИС ЕОП — Регистър на обществените поръчки</h1>
    <nav class="tabs">
      <button id="tab-table" class="tab active" type="button">Регистър</button>
      <button id="tab-stats" class="tab" type="button">Статистика</button>
    </nav>
    <div class="searchbox" id="searchbox">
      <span class="icon">🔍</span>
      <input id="search" type="text" placeholder="Търси в цялата база (номер, възложител, участник, ЕИК, контакт...)" autocomplete="off">
      <button id="searchbtn" type="button">Търси</button>
    </div>
    <div class="meta">
      <span id="count"></span>
      <button id="notifications-btn" class="notif-btn" type="button">Известия</button>
      <button class="refresh" onclick="location.reload()">Обнови</button>
    </div>
  </div>
</header>
<div class="wrap">
  <div id="table-view">
    <div class="toolbar">
    <span class="stats" id="stats"></span>
    <div class="pager">
      <button id="prev">‹ Предишна</button>
      <span class="stats">Страница</span>
      <input id="pageinput" class="pageinput" type="number" min="1" value="1">
      <span class="stats" id="totalpages">от 1</span>
      <button id="goto" class="goto">Към страница</button>
      <button id="next">Следваща ›</button>
      <select id="pagesize">
        <option value="50">50</option>
        <option value="100" selected>100</option>
        <option value="200">200</option>
        <option value="500">500</option>
      </select>
    </div>
  </div>
  <div class="table-scroll">
    <table id="tbl">
      <colgroup id="colgroup"></colgroup>
      <thead id="thead"></thead>
      <tbody id="tbody"></tbody>
    </table>
    <div class="empty" id="empty" style="display:none">Няма намерени записи.</div>
  </div>
  </div>

  <div id="stats-view" style="display:none">
    <h2>Статистика <span class="muted" id="stats-total"></span></h2>
    <div class="stats-grid">
      <div class="stat-card">
        <h3>Топ 10 възложители (най-много поръчки)</h3>
        <div id="top-authorities"></div>
      </div>
      <div class="stat-card">
        <h3>Статуси на изпълнение</h3>
        <div class="pie-wrap">
          <svg id="pie-statuses"></svg>
          <div id="status-legend"></div>
        </div>
      </div>
      <div class="stat-card">
        <h3>Топ 10 най-скъпи поръчки</h3>
        <div id="top-expensive"></div>
        <div class="stats-note">Стойностите са нормализирани към EUR (фиксиран курс 1 EUR = 1.95583 BGN).</div>
      </div>
      <div class="stat-card">
        <h3>Топ 10 поръчки с най-много участници</h3>
        <div id="top-participants"></div>
      </div>
      <div class="stat-card">
        <h3>Поръчки по месеци (януари – декември)</h3>
        <div id="monthly-stats"></div>
      </div>
      <div class="stat-card">
        <h3>Поръчки по години</h3>
        <div id="yearly-stats"></div>
      </div>
      <div class="stat-card">
        <h3>Стойност на поръчките по години</h3>
        <div id="yearly-values"></div>
        <div class="stats-note" id="total-value-note"></div>
        <div class="stats-note">Стойностите са нормализирани към EUR (фиксиран курс 1 EUR = 1.95583 BGN).</div>
      </div>
      <div class="stat-card">
        <h3>Топ 20 градове (най-много поръчки)</h3>
        <div id="top-cities"></div>
        <div class="stats-note">Градът се разпознава от името на възложителя.</div>
      </div>
    </div>
  </div>
</div>

<div id="loading" class="modal" style="display:none">
  <div class="modal-content">
    <div class="spinner"></div>
    <div>Зареждане...</div>
  </div>
</div>

<div id="analysis-modal" class="modal" style="display:none">
  <div class="modal-content analysis-content">
    <div class="analysis-header">
      <span>Анализ и Обобщение на поръчката</span>
      <div style="display:flex;align-items:center;gap:10px">
        <button id="analysis-copy" class="close" type="button" title="Копирай в клипборда" style="font-size:16px;display:none">
          <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="9" y="9" width="13" height="13" rx="2" ry="2"></rect><path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"></path></svg>
        </button>
        <button id="analysis-close" class="close" type="button">✕</button>
      </div>
    </div>
    <div id="analysis-body" class="analysis-body"></div>
    <div class="analysis-actions">
      <button id="analysis-cancel" type="button" style="display:none">Прекрати</button>
      <button id="analysis-ok" type="button" style="display:none">Затвори</button>
    </div>
  </div>
</div>

<div id="notifications-modal" class="modal" style="display:none">
  <div class="modal-content notif-content">
    <div class="analysis-header">
      <span>Известия — условия за имейл при съвпадение</span>
      <button id="notif-close" class="close" type="button">✕</button>
    </div>
    <div class="form-check">
      <input type="checkbox" id="notif-enabled">
      <label for="notif-enabled">Активирано (изпращай имейл при съвпадение)</label>
    </div>
    <div class="form-row">
      <label>Ключови думи (в името или описанието), разделени със запетая</label>
      <input type="text" id="notif-keywords" placeholder="напр. охрана, видеонаблюдение">
      <span class="hint">Съвпада, ако някоя от думите се среща в името ИЛИ описанието.</span>
    </div>
    <div class="form-row">
      <label>Възложител</label>
      <input type="text" id="notif-organization" placeholder="напр. ОБЩИНА СОФИЯ">
      <span class="hint">Част от името на възложителя (без значение на малки/главни букви).</span>
    </div>
    <div class="form-row">
      <label>Обект на поръчката</label>
      <select id="notif-object">
        <option value="">Всички</option>
        <option value="Услуги">Услуги</option>
        <option value="Доставки">Доставки</option>
        <option value="Строителство">Строителство</option>
      </select>
    </div>
    <div class="form-row">
      <label>Имейл адрес(и), разделени със запетая</label>
      <input type="text" id="notif-emails" placeholder="напр. ivan@example.com, petar@example.com">
    </div>
    <div class="detail-title" style="margin-top:10px">SMTP настройки</div>
    <div class="form-row">
      <label>Host</label>
      <input type="text" id="smtp-host" placeholder="smtp.gmail.com">
    </div>
    <div class="form-row">
      <label>Port</label>
      <input type="text" id="smtp-port" placeholder="587">
    </div>
    <div class="form-row">
      <label>Username</label>
      <input type="text" id="smtp-username">
    </div>
    <div class="form-row">
      <label>Password</label>
      <input type="password" id="smtp-password">
    </div>
    <div class="form-row">
      <label>From</label>
      <input type="text" id="smtp-from" placeholder="име@example.com">
    </div>
    <div class="form-check">
      <input type="checkbox" id="smtp-use-tls">
      <label for="smtp-use-tls">Use TLS</label>
    </div>
    <div id="notif-status"></div>
    <div class="form-actions">
      <button id="notif-cancel" type="button">Отказ</button>
      <button id="notif-save" type="button">Запази</button>
    </div>
  </div>
</div>

<script>
const COLUMNS = [
  { key: "special_number", label: "Уникален номер",     type: "text",   pct: 8 },
  { key: "name",           label: "Име",                type: "text",   pct: 20 },
  { key: "organization",   label: "Възложител",         type: "text",   pct: 13 },
  { key: "procedure",      label: "Начин на възлагане", type: "text",   pct: 11 },
  { key: "object",         label: "Обект",              type: "text",   pct: 7 },
  { key: "value",          label: "Прогнозна стойност", type: "number", pct: 9 },
  { key: "date",           label: "Дата на публикуване",type: "date",   pct: 9 },
  { key: "status",         label: "Статус",             type: "text",   pct: 8 },
  { key: "contact",        label: "Лице за контакт",    type: "text",   pct: 11 },
  { key: "participants",   label: "Участници",          type: "number", pct: 4 },
];

const state = { q: "", sort: "date", order: "desc", page: 1, size: 100, total: 0 };
let loadingTimer = null;

function esc(s) {
  return String(s == null ? "" : s)
    .replace(/&/g, "&amp;").replace(/</g, "&lt;").replace(/>/g, "&gt;")
    .replace(/"/g, "&quot;").replace(/'/g, "&#39;");
}

function fmtDate(iso) {
  if (!iso) return "";
  const d = new Date(iso);
  if (isNaN(d)) return esc(iso);
  const p = n => String(n).padStart(2, "0");
  return `${p(d.getDate())}.${p(d.getMonth()+1)}.${d.getFullYear()} ${p(d.getHours())}:${p(d.getMinutes())}`;
}

function fmtNum(v) {
  if (v == null || v === "") return "";
  const n = Number(v);
  if (isNaN(n)) return String(v);
  return n.toLocaleString("bg-BG", { maximumFractionDigits: 2 });
}

function statusInfo(t) {
  const outcome = t.outcome || "";
  const status = t.status || "";
  const label = outcome || status || "—";
  let cls = "neutral";
  if (outcome === "Изпълнена") cls = "ok";
  else if (outcome === "Прекратена" || status === "Прекратена") cls = "term";
  else if (outcome === "Възложена") cls = "award";
  return { label, cls };
}

function buildColgroup() {
  const cg = document.getElementById("colgroup");
  cg.innerHTML = COLUMNS.map(c => `<col style="width:${c.pct}%">`).join("");
  const tbl = document.getElementById("tbl");
  // Таблицата заема цялата ширина и се преоразмерява автоматично с прозореца.
  tbl.style.width = "100%";
  tbl.style.minWidth = "900px";
}

function buildHeader() {
  const thead = document.getElementById("thead");
  let html = "<tr>";
  COLUMNS.forEach((c, i) => {
    const sorted = state.sort === c.key;
    const arrow = sorted ? (state.order === "asc" ? "▲" : "▼") : "↕";
    html += `<th class="sortable${sorted ? " sorted" : ""}" data-key="${c.key}" data-col="${i}" style="position:relative">
      <span class="th-inner"><span class="label">${esc(c.label)}</span><span class="arrow">${arrow}</span></span>
      <span class="resizer" data-col="${i}"></span>
    </th>`;
  });
  html += "</tr>";
  thead.innerHTML = html;
  attachHeaderHandlers();
}

function attachHeaderHandlers() {
  const thead = document.getElementById("thead");
  thead.querySelectorAll("th.sortable").forEach(th => {
    th.addEventListener("click", (e) => {
      if (e.target.classList.contains("resizer")) return;
      const key = th.dataset.key;
      if (state.sort === key) state.order = state.order === "asc" ? "desc" : "asc";
      else { state.sort = key; state.order = "asc"; }
      state.page = 1;
      load();
    });
  });
  // ръчно преоразмеряване на колоните
  setupColumnResize(thead, document.getElementById("colgroup"));
}

function setupColumnResize(theadEl, colgroupEl) {
  theadEl.querySelectorAll(".resizer").forEach(rs => {
    rs.addEventListener("mousedown", (e) => {
      e.preventDefault();
      e.stopPropagation();
      const col = colgroupEl.querySelectorAll("col")[Number(rs.dataset.col)];
      const tbl = theadEl.closest("table");
      const startX = e.pageX;
      const isPct = (col.style.width || "").includes("%");
      const tableW = tbl.getBoundingClientRect().width;
      const startW = isPct
        ? (parseFloat(col.style.width) || (col.getBoundingClientRect().width / tableW * 100))
        : (parseFloat(col.style.width) || col.getBoundingClientRect().width);
      const startTableW = parseFloat(tbl.style.width) || tbl.getBoundingClientRect().width;
      const onMove = (ev) => {
        const delta = ev.pageX - startX;
        if (isPct) {
          const deltaPct = delta / tableW * 100;
          col.style.width = Math.max(3, startW + deltaPct) + "%";
        } else {
          const w = Math.max(40, startW + delta);
          col.style.width = w + "px";
          tbl.style.width = (startTableW + delta) + "px";
          tbl.style.minWidth = tbl.style.width;
        }
      };
      const onUp = () => {
        document.removeEventListener("mousemove", onMove);
        document.removeEventListener("mouseup", onUp);
        document.body.style.cursor = "";
        document.body.style.userSelect = "";
      };
      document.body.style.cursor = "col-resize";
      document.body.style.userSelect = "none";
      document.addEventListener("mousemove", onMove);
      document.addEventListener("mouseup", onUp);
    });
  });
}

function field(label, value) {
  return `<div class="field"><div class="k">${esc(label)}</div><div class="v">${value || '<span class="muted">—</span>'}</div></div>`;
}

function participantsHTML(t) {
  const parts = t.participants || [];
  if (!parts.length) return field("Участници", "");
  let html = `<div class="detail-title">Участници (${parts.length})</div>`;
  html += `<table class="participants"><thead><tr><th>Фирма</th><th>ЕИК / БУЛСТАТ</th><th>Цена</th><th>Papagal</th></tr></thead><tbody>`;
  for (const p of parts) {
    const price = p.price != null ? fmtNum(p.price) : "";
    const cons = p.consortium_member ? ' <span class="chip">член на консорциум</span>' : "";
    const query = (p.registry_number || p.name || "").trim();
    const papagal = query ? `https://papagal.bg/search_results/${encodeURIComponent(query)}` : "";
    html += `<tr><td>${esc(p.name)}${cons}</td><td>${esc(p.registry_number || "")}</td><td>${price}</td>`;
    html += `<td>${papagal
      ? `<a class="papagal-link" href="${esc(papagal)}" target="_blank" rel="noopener" title="Търси в Papagal.bg">🦜</a>`
      : '<span class="muted">—</span>'}</td></tr>`;
  }
  html += `</tbody></table>`;
  return html;
}

function specialtyHTML(t) {
  const s = t.specialty || {};
  const map = { HasGreenCriteria: "Зелени критерии", HasProtectedJobs: "Защитени работни места", HasStrategicRequirements: "Стратегически изисквания", IsEUFunded: "Финансирано от ЕС", IsSecurityAndDefense: "Сигурност и отбрана", IsSocialService: "Социална услуга" };
  const chips = Object.keys(map).filter(k => s[k]).map(k => `<span class="chip">${map[k]}</span>`);
  return chips.length ? `<div class="chips">${chips.join("")}</div>` : '<span class="muted">—</span>';
}

function detailHTML(t) {
  const st = statusInfo(t);
  const est = t.estimated_value != null ? ((t.currency ? t.currency + " " : "") + fmtNum(t.estimated_value)) : "";
  const url = t.tender_id ? `https://app.eop.bg/today/${t.tender_id}` : "";
  let h = '<div class="detail-grid">';
  h += field("Име", esc(t.name));
  h += field("Възложител", esc(t.organization));
  h += field("Уникален номер", esc(t.special_number));
  h += field("Статус", `<span class="badge ${st.cls}">${esc(st.label)}</span>`);
  h += field("Начин на възлагане", esc(t.procedure_type_label));
  h += field("Обект на поръчката", esc(t.type_of_contract_label));
  h += field("Прогнозна стойност", esc(est));
  h += field("Дата на публикуване", fmtDate(t.publication_date));
  h += field("Краен срок", fmtDate(t.deadline));
  h += field("Брой получени оферти", t.offer_count != null ? esc(t.offer_count) : "");
  h += field("Брой обособени позиции", t.lot_count != null ? esc(t.lot_count) : "");
  h += field("Съвместна поръчка", t.is_joint_procurement ? "да" : "не");
  h += field("Основна поръчка", esc(t.parent_tender_special_number));
  h += field("Лице за контакт", esc(t.contact_person));
  h += field("URL", url ? `<a class="link" href="${esc(url)}" target="_blank" rel="noopener">отвори</a>` : "");
  h += `</div>`;
  h += `<div class="detail-title">Особености</div>` + specialtyHTML(t);
  h += participantsHTML(t);
  h += `<button class="analyze-btn" type="button" data-tender-id="${t.tender_id || ''}">Анализ и Обобщение на поръчката</button>`;
  return h;
}

function renderBody(rows) {
  const tbody = document.getElementById("tbody");
  document.getElementById("empty").style.display = rows.length ? "none" : "block";
  let html = "";
  rows.forEach((t, idx) => {
    const st = statusInfo(t);
    const est = t.estimated_value != null ? ((t.currency ? t.currency + " " : "") + fmtNum(t.estimated_value)) : "";
    html += `<tr class="main" data-idx="${idx}">`;
    html += `<td class="num">${esc(t.special_number)}</td>`;
    html += `<td>${esc(t.name)}</td>`;
    html += `<td>${esc(t.organization)}</td>`;
    html += `<td>${esc(t.procedure_type_label)}</td>`;
    html += `<td>${esc(t.type_of_contract_label)}</td>`;
    html += `<td>${esc(est)}</td>`;
    html += `<td>${fmtDate(t.publication_date)}</td>`;
    html += `<td><span class="badge ${st.cls}">${esc(st.label)}</span></td>`;
    html += `<td>${esc(t.contact_person)}</td>`;
    html += `<td>${t.participants_count != null ? esc(t.participants_count) : ""}</td>`;
    html += `</tr>`;
    html += `<tr class="detail" style="display:none"><td colspan="${COLUMNS.length}">${detailHTML(t)}</td></tr>`;
  });
  tbody.innerHTML = html;
}

function attachRowHandlers() {
  const tbody = document.getElementById("tbody");
  tbody.addEventListener("click", (e) => {
    const btn = e.target.closest(".analyze-btn");
    if (btn) {
      analyzeTender(btn.dataset.tenderId);
      return;
    }
    const row = e.target.closest("tr.main");
    if (!row) return;
    const detail = row.nextElementSibling;
    if (!detail || !detail.classList.contains("detail")) return;
    const isOpen = detail.style.display !== "none";
    document.querySelectorAll("tr.detail").forEach(d => d.style.display = "none");
    document.querySelectorAll("tr.main").forEach(r => r.classList.remove("open"));
    if (!isOpen) {
      detail.style.display = "";
      row.classList.add("open");
    }
  });
}

let analysisAbort = null;
let lastAnalysisText = "";

function openAnalysisModal() {
  document.getElementById("analysis-modal").style.display = "flex";
  document.getElementById("analysis-body").textContent = "";
  document.getElementById("analysis-cancel").style.display = "inline-block";
  document.getElementById("analysis-ok").style.display = "none";
  document.getElementById("analysis-copy").style.display = "none";
}

function closeAnalysisModal() {
  if (analysisAbort) { analysisAbort.abort(); analysisAbort = null; }
  document.getElementById("analysis-modal").style.display = "none";
}

function showAnalysisProgress(message) {
  const body = document.getElementById("analysis-body");
  body.style.cursor = "none";
  body.style.overflow = "hidden";
  document.getElementById("analysis-copy").style.display = "none";
  body.innerHTML =
    `<div style="display:flex;align-items:flex-start;gap:12px;color:var(--muted)">
       <div class="spinner" style="flex-shrink:0;margin-top:2px"></div>
       <span style="min-width:0;overflow-wrap:anywhere;word-break:break-word">${esc(message)}</span>
     </div>`;
}

function showAnalysisResult(text) {
  const body = document.getElementById("analysis-body");
  body.style.cursor = "auto";
  body.style.overflow = "auto";
  lastAnalysisText = text || "";
  body.textContent = text || "(няма резултат)";
  document.getElementById("analysis-cancel").style.display = "none";
  document.getElementById("analysis-ok").style.display = "inline-block";
  document.getElementById("analysis-copy").style.display = "inline-block";
}

async function copyAnalysis() {
  const text = lastAnalysisText || document.getElementById("analysis-body").textContent || "";
  if (!text) return;
  const btn = document.getElementById("analysis-copy");
  try {
    if (navigator.clipboard && navigator.clipboard.writeText) {
      await navigator.clipboard.writeText(text);
    } else {
      const ta = document.createElement("textarea");
      ta.value = text;
      document.body.appendChild(ta);
      ta.select();
      document.execCommand("copy");
      document.body.removeChild(ta);
    }
    btn.textContent = "✓";
    btn.style.color = "#1a7f37";
    setTimeout(() => {
      btn.textContent = "";
      btn.innerHTML = '<svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="9" y="9" width="13" height="13" rx="2" ry="2"></rect><path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"></path></svg>';
      btn.style.color = "";
    }, 1500);
  } catch (e) {
    /* ignore */
  }
}

function showAnalysisError(msg) {
  const body = document.getElementById("analysis-body");
  body.style.cursor = "auto";
  body.style.overflow = "auto";
  document.getElementById("analysis-copy").style.display = "none";
  body.innerHTML = `<span style="color:#b42318">${esc(msg)}</span>`;
  document.getElementById("analysis-cancel").style.display = "none";
  document.getElementById("analysis-ok").style.display = "inline-block";
}

async function analyzeTender(tenderId) {
  if (!tenderId) { showAnalysisError("Липсва идентификатор на поръчката."); return; }
  openAnalysisModal();
  showAnalysisProgress("Подготовка на анализа...");
  analysisAbort = new AbortController();
  try {
    const res = await fetch("/api/analyze", {
      method: "POST",
      headers: { "Content-Type": "application/json" },
      body: JSON.stringify({ tender_id: Number(tenderId) }),
      signal: analysisAbort.signal,
    });
    if (!res.ok) {
      showAnalysisError("Грешка при анализа (HTTP " + res.status + ").");
      return;
    }
    const reader = res.body.getReader();
    const decoder = new TextDecoder();
    let buffer = "";
    while (true) {
      const { done, value } = await reader.read();
      if (done) break;
      buffer += decoder.decode(value, { stream: true });
      let nl;
      while ((nl = buffer.indexOf("\n")) >= 0) {
        const line = buffer.slice(0, nl).trim();
        buffer = buffer.slice(nl + 1);
        if (!line) continue;
        let evt;
        try { evt = JSON.parse(line); } catch (e) { continue; }
        if (evt.type === "progress") showAnalysisProgress(evt.message);
        else if (evt.type === "done") showAnalysisResult(evt.analysis);
        else if (evt.type === "error") showAnalysisError(evt.error);
      }
    }
  } catch (err) {
    if (err.name === "AbortError") {
      // прекратено от потребителя — просто затваряме
    } else {
      showAnalysisError("Грешка: " + err.message);
    }
  } finally {
    analysisAbort = null;
  }
}

function showLoading() {
  clearTimeout(loadingTimer);
  loadingTimer = setTimeout(() => {
    document.getElementById("loading").style.display = "flex";
  }, 200);
}

function hideLoading() {
  clearTimeout(loadingTimer);
  document.getElementById("loading").style.display = "none";
}

async function load() {
  showLoading();
  try {
    const params = new URLSearchParams({
      q: state.q, sort: state.sort, order: state.order,
      page: state.page, size: state.size,
    });
    const res = await fetch("/api/tenders?" + params.toString());
    const data = await res.json();
    state.total = data.total || 0;
    const pages = Math.max(1, Math.ceil(state.total / state.size));
    if (state.page > pages) state.page = pages;
    renderBody(data.rows);
    updateSortIndicators();
    const from = state.total ? (state.page - 1) * state.size + 1 : 0;
    const to = Math.min(state.page * state.size, state.total);
    document.getElementById("stats").textContent = `Намерени: ${state.total} поръчки (${from}–${to})`;
    document.getElementById("count").textContent = `${state.total} поръчки`;
    document.getElementById("totalpages").textContent = `от ${pages}`;
    const pageInput = document.getElementById("pageinput");
    pageInput.value = state.page;
    pageInput.max = pages;
    document.getElementById("prev").disabled = state.page <= 1;
    document.getElementById("next").disabled = state.page >= pages;
    document.getElementById("goto").disabled = pages <= 1;
  } catch (err) {
    document.getElementById("stats").textContent = "Грешка при зареждане.";
  } finally {
    hideLoading();
  }
}

function updateSortIndicators() {
  document.querySelectorAll("#thead th").forEach(th => {
    const key = th.dataset.key;
    const sorted = state.sort === key;
    th.classList.toggle("sorted", sorted);
    const arrow = th.querySelector(".arrow");
    if (arrow) arrow.textContent = sorted ? (state.order === "asc" ? "▲" : "▼") : "↕";
  });
}

const STATUS_COLORS = ["#189cff", "#1a7f37", "#f59e0b", "#b42318", "#7c3aed", "#0ea5e9", "#64748b", "#ec4899", "#14b8a6", "#f97316"];

function renderBars(el, items) {
  const max = Math.max(...items.map(i => i.value), 1);
  let html = "";
  items.forEach(it => {
    const w = (it.value / max * 100).toFixed(1);
    html += `<div class="bar-row">
      <div class="bar-label" title="${esc(it.label)}">${esc(it.label)}</div>
      <div class="bar-track"><div class="bar-fill" style="width:${w}%"></div></div>
      <div class="bar-value">${it.value}</div>
    </div>`;
  });
  el.innerHTML = html || '<span class="muted">Няма данни.</span>';
}

function renderYearlyValues(el, items, totalEur) {
  if (!items.length) { el.innerHTML = '<span class="muted">Няма данни.</span>'; return; }
  let html = '<table class="minitable"><thead><tr><th>Година</th><th>Брой</th><th>Обща стойност (EUR)</th></tr></thead><tbody>';
  items.forEach(it => {
    html += `<tr><td><strong>${esc(it.year)}</strong></td><td>${esc(it.cnt)}</td><td>${fmtNum(it.total_eur)}</td></tr>`;
  });
  html += '</tbody></table>';
  el.innerHTML = html;
  const note = document.getElementById("total-value-note");
  if (note) {
    const millions = (totalEur || 0) / 1e6;
    note.textContent = `Общо до момента: ${fmtNum(totalEur || 0)} EUR (≈ ${millions.toFixed(2)} млн. EUR)`;
  }
}

function renderPieChart(svgEl, items) {
  const total = items.reduce((s, i) => s + i.cnt, 0);
  if (!total) { svgEl.innerHTML = ""; return; }
  const cx = 100, cy = 100, r = 90;
  let angle = 0;
  let paths = "";
  items.forEach((it, idx) => {
    const frac = it.cnt / total;
    const a0 = angle, a1 = angle + frac * 2 * Math.PI;
    angle = a1;
    const large = (a1 - a0) > Math.PI ? 1 : 0;
    const x0 = cx + r * Math.cos(a0 - Math.PI / 2);
    const y0 = cy + r * Math.sin(a0 - Math.PI / 2);
    const x1 = cx + r * Math.cos(a1 - Math.PI / 2);
    const y1 = cy + r * Math.sin(a1 - Math.PI / 2);
    paths += `<path d="M ${cx} ${cy} L ${x0} ${y0} A ${r} ${r} 0 ${large} 1 ${x1} ${y1} Z" fill="${STATUS_COLORS[idx % STATUS_COLORS.length]}"><title>${esc(it.label)}: ${it.cnt}</title></path>`;
  });
  svgEl.setAttribute("viewBox", "0 0 200 200");
  svgEl.setAttribute("width", "210");
  svgEl.setAttribute("height", "210");
  svgEl.innerHTML = paths;
}

function renderStatusLegend(el, items) {
  const total = items.reduce((s, i) => s + i.cnt, 0);
  if (!total) { el.innerHTML = '<span class="muted">Няма данни.</span>'; return; }
  let html = '<table class="legend">';
  items.forEach((it, idx) => {
    const pct = (it.cnt / total * 100).toFixed(1);
    html += `<tr>
      <td><span class="dot" style="background:${STATUS_COLORS[idx % STATUS_COLORS.length]}"></span>${esc(it.label)}</td>
      <td><strong>${it.cnt}</strong></td>
      <td class="muted">${pct}%</td>
    </tr>`;
  });
  html += "</table>";
  el.innerHTML = html;
}

function stripTags(s) {
  return String(s == null ? "" : s).replace(/<[^>]+>/g, " ");
}

function renderResizableTable(el, headers, rows, cellFn) {
  if (!rows.length) { el.innerHTML = '<span class="muted">Няма данни.</span>'; return; }

  // Автоматично изчисляване на ширината на колоните според съдържанието.
  const widths = headers.map((h, ci) => {
    let maxLen = h.label.length;
    for (const r of rows) {
      const c = cellFn(r, ci, 0);
      const t = (c.text != null) ? String(c.text) : stripTags(c.html || "");
      maxLen = Math.max(maxLen, t.length);
    }
    return Math.max(60, Math.min(360, Math.round(maxLen * 7.2) + 28));
  });
  const totalW = widths.reduce((a, b) => a + b, 0);

  let html = `<div class="minitable-scroll"><table class="minitable" style="width:${totalW}px;min-width:${totalW}px">`;
  html += `<colgroup>${widths.map(w => `<col style="width:${w}px">`).join("")}</colgroup>`;
  html += `<thead><tr>${headers.map((h, i) =>
    `<th style="position:relative"><span class="th-inner" title="${esc(h.label)}">${esc(h.label)}</span><span class="resizer" data-col="${i}"></span></th>`
  ).join("")}</tr></thead>`;
  html += `<tbody>`;
  rows.forEach((r, ri) => {
    html += `<tr>${headers.map((h, ci) => {
      const c = cellFn(r, ci, ri);
      const title = c.title ? ` title="${esc(c.title)}"` : "";
      return `<td${title}>${c.html}</td>`;
    }).join("")}</tr>`;
  });
  html += `</tbody></table></div>`;
  el.innerHTML = html;
  setupColumnResize(el.querySelector("thead"), el.querySelector("colgroup"));
}

function renderExpensive(el, items) {
  const headers = [
    { label: "#" },
    { label: "Поръчка" },
    { label: "Възложител" },
    { label: "Стойност" },
  ];
  renderResizableTable(el, headers, items, (it, ci, ri) => {
    if (ci === 0) return { html: String(ri + 1), text: String(ri + 1) };
    if (ci === 1) {
      const text = (it.name || "") + " " + (it.special_number || "");
      return { html: `${esc(it.name)}<br><span class="muted">${esc(it.special_number)}</span>`, text, title: it.name };
    }
    if (ci === 2) return { html: esc(it.organization), text: it.organization, title: it.organization };
    if (ci === 3) {
      const val = (it.currency ? it.currency + " " : "") + fmtNum(it.estimated_value);
      return { html: `${esc(val)}<br><span class="muted">≈ EUR ${fmtNum(it.estimated_value_eur)}</span>`, text: val };
    }
    return { html: "" };
  });
}

function renderParticipants(el, items) {
  const headers = [
    { label: "#" },
    { label: "Поръчка" },
    { label: "Възложител" },
    { label: "Участници" },
  ];
  renderResizableTable(el, headers, items, (it, ci, ri) => {
    if (ci === 0) return { html: String(ri + 1), text: String(ri + 1) };
    if (ci === 1) {
      const text = (it.name || "") + " " + (it.special_number || "");
      return { html: `${esc(it.name)}<br><span class="muted">${esc(it.special_number)}</span>`, text, title: it.name };
    }
    if (ci === 2) return { html: esc(it.organization), text: it.organization, title: it.organization };
    if (ci === 3) return { html: `<strong>${it.participants_count}</strong>`, text: String(it.participants_count) };
    return { html: "" };
  });
}

async function loadStats() {
  showLoading();
  try {
    const res = await fetch("/api/stats");
    const data = await res.json();
    renderBars(document.getElementById("top-authorities"), (data.top_authorities || []).map(a => ({ label: a.name, value: a.cnt })));
    renderPieChart(document.getElementById("pie-statuses"), data.statuses || []);
    renderStatusLegend(document.getElementById("status-legend"), data.statuses || []);
    renderExpensive(document.getElementById("top-expensive"), data.top_expensive || []);
    renderParticipants(document.getElementById("top-participants"), data.top_participants || []);
    renderBars(document.getElementById("monthly-stats"), (data.monthly || []).map(m => ({ label: m.label, value: m.cnt })));
    renderBars(document.getElementById("yearly-stats"), (data.yearly || []).map(y => ({ label: String(y.year), value: y.cnt })));
    renderYearlyValues(document.getElementById("yearly-values"), data.yearly_values || [], data.total_value_eur || 0);
    renderBars(document.getElementById("top-cities"), (data.top_cities || []).map(c => ({ label: c.city, value: c.cnt })));
    document.getElementById("stats-total").textContent = data.total ? `(общо ${data.total} поръчки)` : "";
  } catch (err) {
    document.getElementById("top-authorities").textContent = "Грешка при зареждане.";
  } finally {
    hideLoading();
  }
}

function showTab(tab) {
  const isStats = tab === "stats";
  document.getElementById("table-view").style.display = isStats ? "none" : "block";
  document.getElementById("stats-view").style.display = isStats ? "block" : "none";
  document.getElementById("searchbox").style.display = isStats ? "none" : "flex";
  document.getElementById("tab-table").classList.toggle("active", !isStats);
  document.getElementById("tab-stats").classList.toggle("active", isStats);
  if (isStats) loadStats();
}

function toggleNotifFields() {
  const enabled = document.getElementById("notif-enabled").checked;
  ["notif-keywords", "notif-organization", "notif-object", "notif-emails",
   "smtp-host", "smtp-port", "smtp-username", "smtp-password", "smtp-from", "smtp-use-tls"].forEach(id => {
    document.getElementById(id).disabled = !enabled;
  });
}

function openNotificationsModal() {
  document.getElementById("notifications-modal").style.display = "flex";
  document.getElementById("notif-status").textContent = "";
  fetch("/api/notifications")
    .then(r => r.json())
    .then(data => {
      if (data.error) { document.getElementById("notif-status").textContent = "Грешка: " + data.error; return; }
      document.getElementById("notif-enabled").checked = !!data.enabled;
      document.getElementById("notif-keywords").value = data.keywords || "";
      document.getElementById("notif-organization").value = data.organization || "";
      document.getElementById("notif-object").value = data.object || "";
      document.getElementById("notif-emails").value = data.emails || "";
      const smtp = data.smtp || {};
      document.getElementById("smtp-host").value = smtp.host || "";
      document.getElementById("smtp-port").value = smtp.port != null ? smtp.port : 587;
      document.getElementById("smtp-username").value = smtp.username || "";
      document.getElementById("smtp-password").value = smtp.password || "";
      document.getElementById("smtp-from").value = smtp.from || "";
      document.getElementById("smtp-use-tls").checked = smtp.use_tls !== false;
      toggleNotifFields();
    })
    .catch(err => { document.getElementById("notif-status").textContent = "Грешка: " + err.message; });
}

function closeNotificationsModal() {
  document.getElementById("notifications-modal").style.display = "none";
}

async function saveNotifications() {
  const payload = {
    enabled: document.getElementById("notif-enabled").checked,
    keywords: document.getElementById("notif-keywords").value.trim(),
    organization: document.getElementById("notif-organization").value.trim(),
    object: document.getElementById("notif-object").value,
    emails: document.getElementById("notif-emails").value.trim(),
    smtp: {
      host: document.getElementById("smtp-host").value.trim(),
      port: parseInt(document.getElementById("smtp-port").value, 10) || 587,
      username: document.getElementById("smtp-username").value.trim(),
      password: document.getElementById("smtp-password").value,
      from: document.getElementById("smtp-from").value.trim(),
      use_tls: document.getElementById("smtp-use-tls").checked,
    },
  };
  const statusEl = document.getElementById("notif-status");
  statusEl.textContent = "Записване...";
  try {
    const res = await fetch("/api/notifications", {
      method: "POST",
      headers: { "Content-Type": "application/json" },
      body: JSON.stringify(payload),
    });
    const data = await res.json();
    if (!res.ok || data.error) { statusEl.textContent = "Грешка: " + (data.error || res.status); return; }
    statusEl.textContent = "✓ Запазено в web_ui.cfg.";
  } catch (err) {
    statusEl.textContent = "Грешка: " + err.message;
  }
}

function init() {
  buildColgroup();
  buildHeader();
  attachRowHandlers();

  document.getElementById("tab-table").addEventListener("click", () => showTab("table"));
  document.getElementById("tab-stats").addEventListener("click", () => showTab("stats"));

  document.getElementById("analysis-close").addEventListener("click", closeAnalysisModal);
  document.getElementById("analysis-ok").addEventListener("click", closeAnalysisModal);
  document.getElementById("analysis-cancel").addEventListener("click", closeAnalysisModal);
  document.getElementById("analysis-copy").addEventListener("click", copyAnalysis);
  document.getElementById("analysis-modal").addEventListener("click", (e) => {
    if (e.target.id === "analysis-modal") closeAnalysisModal();
  });

  document.getElementById("notifications-btn").addEventListener("click", openNotificationsModal);
  document.getElementById("notif-close").addEventListener("click", closeNotificationsModal);
  document.getElementById("notif-cancel").addEventListener("click", closeNotificationsModal);
  document.getElementById("notif-save").addEventListener("click", saveNotifications);
  document.getElementById("notif-enabled").addEventListener("change", toggleNotifFields);
  document.getElementById("notifications-modal").addEventListener("click", (e) => {
    if (e.target.id === "notifications-modal") closeNotificationsModal();
  });

  const searchInput = document.getElementById("search");
  const runSearch = () => { state.q = searchInput.value.trim(); state.page = 1; load(); };
  document.getElementById("searchbtn").addEventListener("click", runSearch);
  searchInput.addEventListener("keydown", (e) => {
    if (e.key === "Enter") { e.preventDefault(); runSearch(); }
  });

  document.getElementById("prev").addEventListener("click", () => { if (state.page > 1) { state.page--; load(); } });
  document.getElementById("next").addEventListener("click", () => { state.page++; load(); });

  const pageInput = document.getElementById("pageinput");
  const gotoPage = () => {
    const pages = Math.max(1, Math.ceil(state.total / state.size));
    let p = parseInt(pageInput.value, 10);
    if (isNaN(p)) p = 1;
    p = Math.min(Math.max(1, p), pages);
    state.page = p;
    load();
  };
  document.getElementById("goto").addEventListener("click", gotoPage);
  pageInput.addEventListener("keydown", (e) => {
    if (e.key === "Enter") { e.preventDefault(); gotoPage(); }
  });

  const sizeSelect = document.getElementById("pagesize");
  sizeSelect.addEventListener("change", () => {
    state.size = parseInt(sizeSelect.value, 10) || 100;
    state.page = 1;
    load();
  });
  sizeSelect.value = String(state.size);

  load();
}

init();
</script>
</body>
</html>
"""


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"
    db_path = DEFAULT_DB

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/":
            self.serve_html()
        elif path == "/api/tenders":
            self.serve_api(parse_qs(parsed.query))
        elif path == "/api/stats":
            self.serve_stats()
        elif path == "/api/notifications":
            self.serve_notifications_get()
        else:
            self.send_error(404)

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/analyze":
            self.serve_analyze()
        elif parsed.path == "/api/notifications":
            self.serve_notifications_post()
        else:
            self.send_error(404)

    def _read_json_body(self):
        length = int(self.headers.get("Content-Length", 0) or 0)
        raw = self.rfile.read(length) if length else b""
        return json.loads(raw.decode("utf-8") or "{}")

    def serve_notifications_get(self):
        try:
            data = load_notifications_config()
            data["smtp"] = load_smtp_config()
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
            status = 200
        except Exception as exc:
            body = json.dumps({"error": str(exc)}, ensure_ascii=False).encode("utf-8")
            status = 500
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def serve_notifications_post(self):
        try:
            payload = self._read_json_body()
            save_notifications_config(payload)
            body = json.dumps({"ok": True}, ensure_ascii=False).encode("utf-8")
            status = 200
        except Exception as exc:
            body = json.dumps({"error": str(exc)}, ensure_ascii=False).encode("utf-8")
            status = 500
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def serve_analyze(self):
        try:
            length = int(self.headers.get("Content-Length", 0) or 0)
            raw = self.rfile.read(length) if length else b""
            payload = json.loads(raw.decode("utf-8") or "{}")
            tender_id = payload.get("tender_id")
        except Exception:
            tender_id = None

        try:
            self.send_response(200)
            self.send_header("Content-Type", "application/x-ndjson; charset=utf-8")
            self.send_header("Transfer-Encoding", "chunked")
            self.send_header("Cache-Control", "no-cache")
            self.end_headers()
        except Exception:
            return

        def emit(obj):
            data = (json.dumps(obj, ensure_ascii=False) + "\n").encode("utf-8")
            self.wfile.write(f"{len(data):X}\r\n".encode("ascii") + data + b"\r\n")
            self.wfile.flush()

        try:
            if not tender_id:
                emit({"type": "error", "error": "Липсва tender_id в заявката."})
            else:
                for event in analyze_tender_stream(int(tender_id)):
                    emit(event)
        except (BrokenPipeError, ConnectionResetError):
            # Клиентът е прекъснал връзката (бутон „Прекрати“).
            pass
        except Exception as exc:
            try:
                emit({"type": "error", "error": str(exc)})
            except Exception:
                pass
        finally:
            try:
                self.wfile.write(b"0\r\n\r\n")
                self.wfile.flush()
            except Exception:
                pass

    def serve_html(self):
        body = HTML_TEMPLATE.encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def serve_api(self, qs):
        def one(k, default=""):
            v = qs.get(k, [default])[0]
            return v

        try:
            query = one("q", "")
            sort = one("sort", "date")
            order = one("order", "desc")
            page = max(1, int(one("page", "1") or 1))
            size = min(500, max(1, int(one("size", "100") or 100)))
            data = load_page(self.db_path, query, sort, order, page, size)
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        except Exception as exc:
            body = json.dumps({"total": 0, "rows": [], "error": str(exc)}, ensure_ascii=False).encode("utf-8")

        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def serve_stats(self):
        try:
            data = load_stats(self.db_path)
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        except Exception as exc:
            body = json.dumps({"error": str(exc)}, ensure_ascii=False).encode("utf-8")

        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, fmt, *args):
        pass


def main():
    ap = argparse.ArgumentParser(description="Уеб преглед на регистъра на обществените поръчки.")
    ap.add_argument("--port", type=int, default=DEFAULT_PORT)
    ap.add_argument("--db", default=DEFAULT_DB)
    args = ap.parse_args()

    Handler.db_path = args.db
    server = ThreadingHTTPServer((DEFAULT_HOST, args.port), Handler)
    print(f"Уеб интерфейсът е стартиран: http://{DEFAULT_HOST}:{args.port}/")
    print(f"База данни: {args.db}")
    print("Натиснете Ctrl+C за край.\n")

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nСпиране...")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
